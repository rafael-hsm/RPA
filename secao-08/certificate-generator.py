import os

from docx import Document
from docx.shared import Pt

from openpyxl import load_workbook

students = "secao-08\Alunos.xlsx"

table_data = load_workbook(students)

# select 'aba' Nomes
table_sheet = table_data["Nomes"]

for line  in range(2, len(table_sheet["A"]) + 1):
    file_word = Document("secao-08\Certificado1.docx")

    style_doc = file_word.styles["Normal"]

    # Getting name of student
    name_student = table_sheet['A%s' % line].value

    for i in file_word.paragraphs:

        if "@nome" in i.text:
            i.text = name_student
            font = style_doc.font
            font.name = "Calibri (Corpo)"
            font.size = Pt(24)


    parth_certificate = f"C:\\Users\\EBJB\\Cursos\\RPA\\secao-08\\certificates\\" + name_student + ".docx" 

    file_word.save(parth_certificate)

print("Certificate gerated successfully!")
